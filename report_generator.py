"""
Report Generation and Export Module for Student Performance Analysis System.
Produces formatted CSV, multi-sheet Excel reports with OpenPyXL, and comprehensive
academic summaries suitable for administration and faculty reviews.
"""

from io import BytesIO
from datetime import datetime
from typing import Dict, Any, Optional
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from analytics import enrich_academic_records, calculate_kpis, get_subject_performance, get_grade_distribution


def export_dataframe_to_csv(df: pd.DataFrame) -> bytes:
    """Export a sanitized DataFrame to CSV bytes."""
    # Ensure sensitive columns are excluded if any
    safe_df = df.copy()
    for col in ["password_hash", "internal_student_id"]:
        if col in safe_df.columns:
            safe_df.drop(columns=[col], inplace=True)
    return safe_df.to_csv(index=False).encode("utf-8")


def generate_excel_academic_report(
    records_df: pd.DataFrame,
    students_df: pd.DataFrame,
    report_title: str = "Student Performance Comprehensive Report",
    filter_metadata: Optional[Dict[str, Any]] = None,
) -> bytes:
    """
    Generate a polished multi-tab Excel workbook with summary metrics,
    subject breakdowns, at-risk flags, and full detailed records.
    """
    output = BytesIO()
    enriched = enrich_academic_records(records_df)

    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        # Sheet 1: Executive Summary
        kpis = calculate_kpis(records_df, len(students_df))
        summary_rows = [
            {"Metric": "Report Title", "Value": report_title},
            {"Metric": "Generated On", "Value": datetime.now().strftime("%Y-%m-%d %H:%M:%S")},
            {"Metric": "Department Filter", "Value": filter_metadata.get("department", "All") if filter_metadata else "All"},
            {"Metric": "Semester Filter", "Value": filter_metadata.get("semester", "All") if filter_metadata else "All"},
            {"Metric": "Subject Filter", "Value": filter_metadata.get("subject", "All") if filter_metadata else "All"},
            {"Metric": "--------------------------", "Value": "--------------------------"},
            {"Metric": "Total Registered Students", "Value": kpis["total_students"]},
            {"Metric": "Total Academic Assessments", "Value": kpis["total_records"]},
            {"Metric": "Average Total Marks", "Value": f"{kpis['avg_marks']} / 100"},
            {"Metric": "Highest Total Marks", "Value": kpis["highest_marks"]},
            {"Metric": "Lowest Total Marks", "Value": kpis["lowest_marks"]},
            {"Metric": "Overall Pass Percentage", "Value": f"{kpis['pass_percentage']}%"},
            {"Metric": "Passed Assessments", "Value": kpis["pass_count"]},
            {"Metric": "Failed Assessments", "Value": kpis["fail_count"]},
            {"Metric": "Average Attendance", "Value": f"{kpis['avg_attendance']}%"},
            {"Metric": "Students Requiring Support", "Value": kpis["support_needed_count"]},
        ]
        summary_df = pd.DataFrame(summary_rows)
        summary_df.to_excel(writer, sheet_name="Executive Summary", index=False)

        # Sheet 2: Subject Performance Breakdown
        if not records_df.empty:
            subj_perf = get_subject_performance(records_df)
            subj_perf.to_excel(writer, sheet_name="Subject Breakdown", index=False)

        # Sheet 3: Grade Distribution
        if not records_df.empty:
            grade_dist = get_grade_distribution(records_df)
            grade_dist.to_excel(writer, sheet_name="Grade Distribution", index=False)

        # Sheet 4: Detailed Academic Records
        if not enriched.empty:
            columns_export = [
                "student_id",
                "full_name",
                "department",
                "course",
                "semester",
                "subject",
                "internal_marks",
                "assignment_marks",
                "exam_marks",
                "total_marks",
                "grade",
                "grade_desc",
                "attendance_percentage",
                "is_pass",
                "exam_date",
            ]
            cols = [c for c in columns_export if c in enriched.columns]
            export_details = enriched[cols].copy()
            export_details.rename(
                columns={
                    "student_id": "Student ID",
                    "full_name": "Full Name",
                    "department": "Department",
                    "course": "Course",
                    "semester": "Semester",
                    "subject": "Subject",
                    "internal_marks": "Internal (Max 20)",
                    "assignment_marks": "Assignment (Max 10)",
                    "exam_marks": "Exam (Max 70)",
                    "total_marks": "Total (100)",
                    "grade": "Grade",
                    "grade_desc": "Grade Remarks",
                    "attendance_percentage": "Attendance %",
                    "is_pass": "Passed",
                    "exam_date": "Exam Date",
                },
                inplace=True,
            )
            export_details.to_excel(writer, sheet_name="Academic Records", index=False)

        # Format worksheets
        wb = writer.book
        header_fill = PatternFill(start_color="1E88E5", end_color="1E88E5", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        thin_border = Border(
            left=Side(style="thin", color="D0D7DE"),
            right=Side(style="thin", color="D0D7DE"),
            top=Side(style="thin", color="D0D7DE"),
            bottom=Side(style="thin", color="D0D7DE"),
        )

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            for col_idx in range(1, ws.max_column + 1):
                cell = ws.cell(row=1, column=col_idx)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")

            # Auto-adjust column widths
            for col in ws.columns:
                max_len = 0
                col_letter = col[0].column_letter
                for cell in col:
                    if cell.row != 1:
                        cell.border = thin_border
                    val_str = str(cell.value or "")
                    if len(val_str) > max_len:
                        max_len = len(val_str)
                ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    output.seek(0)
    return output.getvalue()


def generate_student_transcript_text(
    student_info: Dict[str, Any],
    records_df: pd.DataFrame,
) -> str:
    """Generate a clean ASCII / markdown transcript card for a single student."""
    enriched = enrich_academic_records(records_df)

    lines = []
    lines.append("=" * 72)
    lines.append("       OFFICIAL STUDENT PERFORMANCE TRANSCRIPT REPORT")
    lines.append("=" * 72)
    lines.append(f"Student ID     : {student_info.get('student_id')}")
    lines.append(f"Student Name   : {student_info.get('full_name')}")
    lines.append(f"Email          : {student_info.get('email')}")
    lines.append(f"Department     : {student_info.get('department')}")
    lines.append(f"Course         : {student_info.get('course')} (Semester {student_info.get('semester')})")
    lines.append(f"Enrollment Date: {student_info.get('enrollment_date')}")
    lines.append(f"Generated On   : {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    lines.append("-" * 72)
    lines.append(f"{'Subject':<24} | {'Int':>4} | {'Asgn':>4} | {'Exam':>4} | {'Total':>5} | {'Grd':>3} | {'Att%':>5}")
    lines.append("-" * 72)

    total_marks_accum = 0.0
    total_attendance_accum = 0.0
    passed_subjects = 0

    if enriched.empty:
        lines.append("No academic records recorded yet.")
    else:
        for _, row in enriched.iterrows():
            total_marks_accum += row["total_marks"]
            total_attendance_accum += row["attendance_percentage"]
            if row["is_pass"]:
                passed_subjects += 1
            lines.append(
                f"{row['subject'][:24]:<24} | {row['internal_marks']:4.1f} | {row['assignment_marks']:4.1f} | "
                f"{row['exam_marks']:4.1f} | {row['total_marks']:5.1f} | {row['grade']:>3} | {row['attendance_percentage']:4.1f}%"
            )

        n = len(enriched)
        avg_score = round(total_marks_accum / n, 2)
        avg_att = round(total_attendance_accum / n, 1)

        lines.append("-" * 72)
        lines.append(f"Summary Statistics for {n} Subject(s):")
        lines.append(f"  Overall Average Score : {avg_score} / 100")
        lines.append(f"  Average Attendance    : {avg_att}%")
        lines.append(f"  Subjects Cleared      : {passed_subjects} / {n}")
        status = "Good Standing" if avg_score >= 50 and avg_att >= 75 else "Academic Support Recommended"
        lines.append(f"  Academic Standing     : {status}")

    lines.append("=" * 72)
    return "\n".join(lines)
